#!/usr/bin/env python3
"""Build an Excel workbook of the daily forecasts: actual price against P10/P50/P90, half-hour by half-hour.

Usage: scripts/build_daily_forecasts_xlsx.py [--region VIC1] [--results results] [--out results/VIC1/daily_forecasts.xlsx]
       (also: make report)

Reads results/<region>/predictions_<model>.parquet (from `make backtest`) and writes four sheets:
  About         what the workbook is, how to use it, definitions and assumptions
  Day viewer    pick a forecast day and see its 48 half-hours and two charts
  Days          one row per forecast day: actual price range and each model's error
  Forecasts     every half-hour: actual, each model's P10/P50/P90 and its absolute error

Values are written as data; every derived number (errors, daily summaries, the viewer) is an Excel formula, so the
workbook recalculates if you edit it. Needs pandas, pyarrow and openpyxl (`make setup` installs them).

The file is written without calculated values: Excel (and LibreOffice) calculate every formula when it is opened.
Viewers that do not calculate, such as some previewers, show the formula cells empty.
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STEPS = 48
MODELS = [  # (saved predictions name, label, line colour)
    ("lightgbm", "LightGBM", "D97706"),
    ("chronos2_zs", "Chronos-2 zero-shot", "2B6CB0"),
    ("chronos2_ft", "Chronos-2 fine-tuned", "2F855A"),
]
FONT = "Arial"
INK, MUTED, HEADER_FILL, INPUT_FILL = "1F2933", "616E7C", "1F2933", "FFF3B0"
PRICE, DATE, DATETIME, TIME = "#,##0.00;-#,##0.00", "yyyy-mm-dd", "yyyy-mm-dd hh:mm", "hh:mm"


def load(results: Path, region: str) -> pd.DataFrame:
    """One row per forecast day and half-hour: actual plus every model's p10, p50, p90."""
    frames = []
    for name, _, _ in MODELS:
        path = results / region / f"predictions_{name}.parquet"
        if not path.is_file():
            raise SystemExit(f"{path} not found: run `make backtest` first")
        df = pd.read_parquet(path).sort_values(["origin", "step"]).reset_index(drop=True)
        frames.append(df.rename(columns={c: f"{name}_{c}" for c in ("p10", "p50", "p90")}))
    base = frames[0][["origin", "step", "timestamp", "actual"]].copy()
    for (name, _, _), df in zip(MODELS, frames):
        if not (df[["origin", "step"]].equals(base[["origin", "step"]]) and np.allclose(df["actual"], base["actual"])):
            raise SystemExit(f"{name} was not scored on the same days and actual prices as {MODELS[0][0]}")
        for c in ("p10", "p50", "p90"):
            base[f"{name}_{c}"] = df[f"{name}_{c}"].to_numpy()
    if not (base.groupby("origin").size() == STEPS).all():
        raise SystemExit(f"every forecast day must have exactly {STEPS} half-hours")
    return base


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_forecasts(ws, df: pd.DataFrame) -> int:
    headers = ["Forecast day", "Half-hour start", "Time", "Actual"]
    for _, label, _ in MODELS:
        headers += [f"{label} P10", f"{label} P50", f"{label} P90"]
    headers += [f"{label} abs error" for _, label, _ in MODELS]
    ws.append(headers)
    formats = [DATE, DATETIME, TIME] + [PRICE] * (len(headers) - 3)
    for i, row in enumerate(df.itertuples(index=False), start=2):
        values = [row.origin.date(), row.timestamp.to_pydatetime(), row.timestamp.time(), float(row.actual)]
        for name, _, _ in MODELS:
            values += [float(getattr(row, f"{name}_{c}")) for c in ("p10", "p50", "p90")]
        # absolute error of each model's median: |actual - P50|; P50 columns are F, I and L
        values += [f"=ABS($D{i}-{col}{i})" for col in ("F", "I", "L")]
        ws.append(values)
    last = len(df) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell, fmt in zip(row, formats):
            cell.number_format = fmt
            cell.font = Font(name=FONT, size=10)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 32
    for c, w in enumerate([13, 17, 8, 10] + [11] * 9 + [12] * 3, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    # shade actual prices that fall outside the fine-tuned model's P10-P90 band
    ws.conditional_formatting.add(
        f"D2:D{last}", FormulaRule(formula=["OR($D2<$K2,$D2>$M2)"], fill=PatternFill("solid", bgColor="F8D7DA"))
    )
    return last


def build_days(ws, n_days: int, last: int) -> int:
    """One row per forecast day. Row n covers Forecasts rows (n-1)*K+1 to n*K of the data, where K is in cell K1."""
    ws["J1"], ws["K1"] = "Half-hours per day", STEPS
    ws["J1"].font = Font(name=FONT, size=10, color=MUTED)
    ws["K1"].font = Font(name=FONT, size=10, bold=True)
    ws["J2"] = "Fixed by the forecast set-up (48 x 30 minutes). The formulas rely on each day having exactly this many rows."
    ws["J2"].font = Font(name=FONT, size=9, color=MUTED)
    ws["J2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("J2:K5")
    headers = ["Day no.", "Forecast day", "Actual average", "Actual max", "Actual min"]
    headers += [f"{label} MAE" for _, label, _ in MODELS] + ["Best model (lowest MAE)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    column = lambda col: f"Forecasts!${col}$2:${col}${last}"  # noqa: E731 - a whole data column
    for n in range(1, n_days + 1):
        r = n + 1
        span = lambda col: (  # noqa: E731, B023 - the 48 rows of this day in one column
            f"INDEX({column(col)},($A{r}-1)*$K$1+1):INDEX({column(col)},$A{r}*$K$1)"
        )
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=f"=INDEX({column('A')},($A{r}-1)*$K$1+1)")
        ws.cell(row=r, column=3, value=f"=AVERAGE({span('D')})")
        ws.cell(row=r, column=4, value=f"=MAX({span('D')})")
        ws.cell(row=r, column=5, value=f"=MIN({span('D')})")
        for k, col in enumerate(("N", "O", "P")):
            ws.cell(row=r, column=6 + k, value=f"=AVERAGE({span(col)})")
        names = ",".join(f'"{label}"' for _, label, _ in MODELS)
        ws.cell(row=r, column=9, value=f"=CHOOSE(MATCH(MIN(F{r}:H{r}),F{r}:H{r},0),{names})")
        for c, fmt in zip(range(1, 10), [ "0", DATE, PRICE, PRICE, PRICE, PRICE, PRICE, PRICE, "General"]):
            ws.cell(row=r, column=c).number_format = fmt
            ws.cell(row=r, column=c).font = Font(name=FONT, size=10)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 32
    for c, w in enumerate([8, 13, 13, 12, 12, 13, 13, 13, 24, 22, 8], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{n_days + 1}"
    return n_days + 1


def build_viewer(ws, n_days: int, last: int, days_last: int, default_day) -> None:
    ws["A1"] = "Daily forecast viewer"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
    ws["A3"], ws["B3"] = "Forecast day (choose)", default_day
    ws["B3"].number_format = DATE
    ws["B3"].fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws["B3"].border = Border(*(Side(style="thin", color=MUTED),) * 4)
    ws["C3"] = "Yellow cell: pick any date from the list."
    ws["A4"], ws["B4"] = "Day no.", f"=MATCH($B$3,Days!$B$2:$B${days_last},0)"
    dv = DataValidation(type="list", formula1=f"=Days!$B$2:$B${days_last}", allow_blank=False,
                        showErrorMessage=True, errorTitle="Forecast day", error="Choose a date from the list.")
    ws.add_data_validation(dv)
    dv.add("B3")
    summary = [("Actual average ($/MWh)", "C"), ("Actual max ($/MWh)", "D"), ("Actual min ($/MWh)", "E")]
    summary += [(f"{label}: mean absolute error", col) for (_, label, _), col in zip(MODELS, "FGH")]
    summary += [("Best model for this day", "I")]
    for i, (label, col) in enumerate(summary, start=6):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=f"=INDEX(Days!${col}$2:${col}${days_last},$B$4)")
        cell.number_format = PRICE if col != "I" else "General"
        cell.alignment = Alignment(horizontal="right")
    first, bottom = 15, 15 + STEPS - 1
    headers = ["Time", "Actual", "LightGBM P50", "Chronos-2 zero-shot P50", "Chronos-2 fine-tuned P10",
               "Chronos-2 fine-tuned P50", "Chronos-2 fine-tuned P90"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=first - 1, column=c, value=h)
    # Forecasts columns: time C, actual D, LightGBM P50 F, zero-shot P50 I, fine-tuned P10 K, P50 L, P90 M
    sources = ["C", "D", "F", "I", "K", "L", "M"]
    for k in range(STEPS):
        r = first + k
        for c, col in enumerate(sources, start=1):
            cell = ws.cell(row=r, column=c,
                           value=f"=INDEX(Forecasts!${col}$2:${col}${last},($B$4-1)*Days!$K$1+ROWS(A${first}:A{r}))")
            cell.number_format = TIME if c == 1 else PRICE
    for row in ws.iter_rows(min_row=3, max_row=bottom, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
    ws["C3"].font = Font(name=FONT, size=9, color=MUTED)
    style_header(ws, first - 1, len(headers))
    ws.row_dimensions[first - 1].height = 44
    for c, w in enumerate([24, 13, 13, 15, 15, 15, 15], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    cats = Reference(ws, min_col=1, min_row=first, max_row=bottom)

    def line_chart(title: str, series: list[tuple[int, str, str, float, str | None]], anchor: str) -> None:
        ch = LineChart()
        ch.title, ch.height, ch.width = f"{title} ($/MWh, NEM time)", 10, 22
        ch.y_axis.number_format = "#,##0"
        ch.x_axis.number_format, ch.x_axis.tickLblSkip = "hh:mm", 4
        ch.x_axis.tickLblPos = "low"  # time labels along the bottom, not on the zero line through the plot
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.legend.position = "b"
        ch.legend.overlay = False  # reserve space for the legend instead of drawing over the time labels
        for col, name, color, width, dash in series:
            ch.add_data(Reference(ws, min_col=col, min_row=first - 1, max_row=bottom), titles_from_data=True)
            s = ch.series[-1]
            s.smooth, s.marker.symbol = False, "none"
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = int(width * 12700)
            if dash:
                s.graphicalProperties.line.dashStyle = dash
        ch.set_categories(cats)
        ws.add_chart(ch, anchor)

    line_chart("Actual price and each model's median (P50)",
               [(2, "Actual", INK, 2.5, None), (3, "LightGBM", MODELS[0][2], 1.75, None),
                (4, "Zero-shot", MODELS[1][2], 1.75, None), (6, "Fine-tuned", MODELS[2][2], 1.75, None)], "I3")
    line_chart("Fine-tuned Chronos-2: P10, P50 and P90 band against the actual price",
               [(2, "Actual", INK, 2.5, None), (5, "P10", MODELS[2][2], 1.25, "dash"),
                (6, "P50", MODELS[2][2], 2.0, None), (7, "P90", MODELS[2][2], 1.25, "dash")], "I23")


def build_about(ws, n_days: int, first_day: str, last_day: str, region: str) -> None:
    rows = [
        ("title", f"{region} daily price forecasts: actual price against P10, P50 and P90"),
        ("text", f"Half-hour forecasts of the Victorian electricity price, made every midnight (NEM time) for the next "
                 f"24 hours, for {n_days} days from {first_day} to {last_day}, with the actual price beside them."),
        ("blank", ""),
        ("head", "Sheets"),
        ("kv", ("Day viewer", "Pick a forecast day (yellow cell) and see its 48 half-hours and two charts.")),
        ("kv", ("Days", "One row per day: the actual price range and each model's error. Filter or sort it to find days.")),
        ("kv", ("Forecasts", "Every half-hour: actual, each model's P10, P50 and P90, and the absolute error of the median.")),
        ("blank", ""),
        ("head", "How to use it"),
        ("text", "1. Open Day viewer and choose a date in the yellow cell. The table and both charts update."),
        ("text", "2. To find an interesting day, use the filter on Days (for example the largest 'Actual max', or the "
                 "largest MAE for a model), then pick that date in Day viewer."),
        ("text", "3. On Forecasts, actual prices shaded red fall outside the fine-tuned model's P10-P90 band."),
        ("blank", ""),
        ("head", "Definitions"),
        ("kv", ("Actual", "The 30-minute average of AEMO's 5-minute dispatch price (RRP) for VIC1, in $/MWh.")),
        ("kv", ("P10, P50, P90", "The model's 10th, 50th (median) and 90th percentile forecast. The actual price should fall "
                                  "below P10 about 10% of the time and above P90 about 10% of the time.")),
        ("kv", ("Absolute error", "|Actual - P50|, in $/MWh.")),
        ("kv", ("MAE", "Mean absolute error: the average of the 48 absolute errors of a day.")),
        ("kv", ("Times", "NEM time (AEST, UTC+10, no daylight saving). A half-hour is labelled by when it starts.")),
        ("blank", ""),
        ("head", "The models"),
        ("kv", ("LightGBM", "Gradient-boosted trees, trained once on 2023-24 prices (lags, recent price statistics, calendar).")),
        ("kv", ("Chronos-2 zero-shot", "Amazon's pretrained Chronos-2 model, not trained on this data; reads the previous 14 days.")),
        ("kv", ("Chronos-2 fine-tuned", "The same model fine-tuned (LoRA) on 2023-24 prices only; reads the previous 14 days.")),
        ("text", "No model sees the forecast day itself; the models that learn from data never trained on 2025 or later."),
        ("blank", ""),
        ("head", "Assumptions"),
        ("text", "Each day has exactly 48 half-hours and the Forecasts sheet is sorted by day and time. Days and Day viewer "
                 "find a day's rows by position using that fact (the 48 is in Days!K1). Do not re-sort the Forecasts sheet."),
        ("text", "Days where the actual prices had gaps (28 days) are not in the forecast set."),
        ("blank", ""),
        ("text", "Source: AEMO public NEM data (VIC1 dispatch price), (c) Australian Energy Market Operator Limited. "
                 "Forecasts: github.com/janxl/nem-timeseries (rebuild this file with `make report`)."),
    ]
    r = 1
    for kind, content in rows:
        if kind == "kv":
            ws.cell(row=r, column=1, value=content[0]).font = Font(name=FONT, size=10, bold=True)
            c = ws.cell(row=r, column=2, value=content[1])
            c.font, c.alignment = Font(name=FONT, size=10), Alignment(wrap_text=True, vertical="top")
        elif kind in ("title", "head", "text"):
            c = ws.cell(row=r, column=1, value=content)
            c.font = {"title": Font(name=FONT, size=14, bold=True, color=INK),
                      "head": Font(name=FONT, size=11, bold=True, color=INK),
                      "text": Font(name=FONT, size=10, color=INK)}[kind]
            if kind == "text":
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = 15 * (1 + len(content) // 105)
        r += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 95
    for r_ in range(1, r):
        if ws.cell(row=r_, column=2).value and len(str(ws.cell(row=r_, column=2).value)) > 95:
            ws.row_dimensions[r_].height = 15 * (1 + len(str(ws.cell(row=r_, column=2).value)) // 95)
    ws.sheet_view.showGridLines = False


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--region", default="VIC1")
    ap.add_argument("--results", type=Path, default=Path("results"))
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--days", type=int, default=None, help="only the first N forecast days (for quick tests)")
    args = ap.parse_args()
    out = args.out or args.results / args.region / "daily_forecasts.xlsx"

    df = load(args.results, args.region)
    if args.days:
        keep = df["origin"].drop_duplicates().iloc[: args.days]
        df = df[df["origin"].isin(keep)].reset_index(drop=True)
    days = df["origin"].drop_duplicates().reset_index(drop=True)
    # open the viewer on a typical day: the fine-tuned model's error closest to its median
    err = (df["actual"] - df["chronos2_ft_p50"]).abs().groupby(df["origin"]).mean()
    default_day = (err - err.median()).abs().idxmin().date()

    wb = Workbook()
    about, viewer, dws, fws = wb.active, wb.create_sheet("Day viewer"), wb.create_sheet("Days"), wb.create_sheet("Forecasts")
    about.title = "About"
    last = build_forecasts(fws, df)
    days_last = build_days(dws, len(days), last)
    build_viewer(viewer, len(days), last, days_last, default_day)
    build_about(about, len(days), str(days.iloc[0].date()), str(days.iloc[-1].date()), args.region)
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when the file is opened
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(days)} days, {len(df)} half-hours, viewer opens on {default_day}")


if __name__ == "__main__":
    main()
